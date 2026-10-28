from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt
from docx.styles.style import BaseStyle
from docx.table import _Row

from aitrace.report.i18n import DEFAULT_LANG, t
from aitrace.report.model import (
    COLUMNS,
    SECTION_KEYS,
    TIMELINE_LIMIT_PDF_DOCX,
    ReportModel,
)

EAST_ASIAN_FONT = "Malgun Gothic"
MONO_FONT = "Consolas"
TABLE_FONT_PT = 9
MONO_FONT_PT = 7.5
PAGE_MARGIN_IN = 0.8
_STYLES_WITH_FONT: tuple[str, ...] = ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet")
_MONO_KEYS = frozenset({"sha256", "media_sha256", "artifact_id", "id"})
_MONO_LABEL_KEYS = ("label.media_sha256", "label.acquisition_hashes", "label.evidence_id", "label.case_id")

_CoverageCol = tuple[str, str, str, int]
_COVERAGE_COLUMNS: dict[str, list[_CoverageCol]] = {
    "plugin_runs": [
        ("name", "플러그인", "Plugin", 3),
        ("version", "버전", "Version", 2),
        ("status", "상태", "Status", 2),
        ("records", "레코드", "Records", 1),
        ("warnings", "경고", "Warnings", 1),
        ("error", "오류", "Error", 4),
    ],
    "parse_warnings": [
        ("plugin", "플러그인", "Plugin", 2),
        ("message", "메시지", "Message", 5),
        ("source_path", "원본 경로", "Source path", 5),
    ],
    "not_decoded": [
        ("tool", "도구", "Tool", 2),
        ("status", "상태", "Status", 3),
        ("detail", "내용", "Detail", 4),
        ("source_path", "원본 경로", "Source path", 5),
    ],
    "findings": [
        ("severity_label", "중요도", "Severity", 1),
        ("title", "제목", "Title", 4),
        ("narrative", "상세", "Narrative", 8),
    ],
}


def _set_east_asian_font(style: BaseStyle) -> None:
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), EAST_ASIAN_FONT)


def _repeat_header(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Mapping):
        return "; ".join(f"{k}: {_cell(v)}" for k, v in value.items())
    if isinstance(value, (list, tuple)):
        return "\n".join(_cell(v) for v in value)
    return str(value)


def _style_run(run: Any, *, mono: bool = False, bold: bool = False) -> None:
    if mono:
        run.font.name = MONO_FONT
        run.font.size = Pt(MONO_FONT_PT)
    else:
        run.font.size = Pt(TABLE_FONT_PT)
    if bold:
        run.bold = True


def _pair(doc: DocumentObject, label: str, value: object, *, mono: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    run = paragraph.add_run(_cell(value))
    if mono:
        run.font.name = MONO_FONT
        run.font.size = Pt(MONO_FONT_PT)


def _pairs(doc: DocumentObject, pairs: Sequence[tuple[str, str]], lang: str) -> None:
    mono_labels = {t(key, lang) for key in _MONO_LABEL_KEYS}
    for label, value in pairs:
        _pair(doc, label, value, mono=label in mono_labels)


def _bullets(doc: DocumentObject, items: Sequence[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _usable_width(doc: DocumentObject) -> int:
    section = doc.sections[-1]
    return int((section.page_width or 0) - (section.left_margin or 0) - (section.right_margin or 0))


def _table(
    doc: DocumentObject,
    columns: Sequence[tuple[str, str, int, bool]],
    rows: Sequence[Sequence[object]],
    empty_text: str,
) -> None:
    if not rows:
        doc.add_paragraph(empty_text)
        return
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    total = sum(weight for _, _, weight, _ in columns)
    usable = _usable_width(doc)
    widths = [Emu(int(usable * weight / total)) for _, _, weight, _ in columns]
    for col, width in zip(table.columns, widths, strict=True):
        col.width = width
    header = table.rows[0]
    _repeat_header(header)
    for cell, (_, label, _, _), width in zip(header.cells, columns, widths, strict=True):
        cell.width = width
        _style_run(cell.paragraphs[0].add_run(label), bold=True)
    for row in rows:
        for cell, (_, _, _, mono), width, value in zip(
            table.add_row().cells, columns, widths, row, strict=True
        ):
            cell.width = width
            _style_run(cell.paragraphs[0].add_run(_cell(value)), mono=mono)


def _spec_table(doc: DocumentObject, section: str, rows: Sequence[Mapping[str, Any]], lang: str) -> None:
    spec = COLUMNS[section]
    columns = [(key, t(label_key, lang), weight, key in _MONO_KEYS) for key, label_key, weight in spec]
    body = [[row.get(key) for key, _, _ in spec] for row in rows]
    _table(doc, columns, body, t("summary.none", lang))


def _write_case(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _pairs(doc, model.case_pairs(), lang)


def _write_evidence(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    groups = model.evidence_pairs()
    if not groups:
        doc.add_paragraph(t("summary.none", lang))
    for i, pairs in enumerate(groups):
        if i:
            doc.add_paragraph()
        _pairs(doc, pairs, lang)


def _write_tools(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _pairs(doc, model.tools_pairs(), lang)


def _write_methodology(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    for paragraph in model.methodology["paragraphs"]:
        doc.add_paragraph(paragraph)
    doc.add_paragraph().add_run(
        model.methodology.get("standards_intro") or t("methodology.standards_intro", lang)
    ).bold = True
    _bullets(doc, model.methodology["standards"])
    _pair(doc, t("label.tool_validation", lang), model.methodology["tool_validation"])


def _write_summary(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    for item in model.summary:
        doc.add_paragraph().add_run(item["question"]).bold = True
        _bullets(doc, item["bullets"])


def _write_ai_tools(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "ai_tools", model.ai_tools, lang)


def _write_ai_queries(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "ai_queries", model.ai_queries, lang)


def _write_file_lifecycle(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "file_lifecycle", model.file_lifecycle, lang)


def _write_exfil(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "exfil", model.exfil, lang)


def _write_attribution(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "attribution", model.attribution, lang)


def _write_timeline(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "timeline", model.timeline_render(TIMELINE_LIMIT_PDF_DOCX), lang)
    note = model.timeline_note(TIMELINE_LIMIT_PDF_DOCX, lang)
    if note is not None:
        doc.add_paragraph(note)


def _write_exports(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _spec_table(doc, "exports", model.exports, lang)


def _write_coverage(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    for name, label in (
        ("plugin_runs", "coverage.plugin_runs"),
        ("parse_warnings", "coverage.parse_warnings"),
        ("not_decoded", "coverage.not_decoded"),
        ("findings", "coverage.findings"),
    ):
        doc.add_paragraph().add_run(t(label, lang)).bold = True
        spec = _COVERAGE_COLUMNS[name]
        columns = [
            (key, ko if lang == "ko" else en, weight, key in _MONO_KEYS) for key, ko, en, weight in spec
        ]
        rows: Sequence[Mapping[str, Any]] = model.coverage.get(name, [])
        body = [[row.get(key) for key, _, _, _ in spec] for row in rows]
        _table(doc, columns, body, t("summary.none", lang))


def _write_limitations(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    _bullets(doc, model.limitations)


def _write_custody(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    columns = [
        ("seq", t("custody.seq", lang), 1, False),
        ("actor", t("custody.actor", lang), 2, False),
        ("action", t("custody.action", lang), 2, False),
        ("detail", t("custody.detail", lang), 6, False),
    ]
    rows = [[e["seq"], e["actor"], e["action"], e["detail"]] for e in model.custody]
    _table(doc, columns, rows, t("summary.none", lang))


def _write_signature(doc: DocumentObject, model: ReportModel, lang: str) -> None:
    columns = [
        ("examiner", t("signature.examiner", lang), 3, False),
        ("signature", t("signature.signature", lang), 4, False),
        ("date", t("signature.date", lang), 2, False),
    ]
    _table(doc, columns, [[model.signature.get("examiner"), "", ""]], t("summary.none", lang))


_SECTION_WRITERS: dict[str, Callable[[DocumentObject, ReportModel, str], None]] = {
    "case": _write_case,
    "evidence": _write_evidence,
    "tools": _write_tools,
    "methodology": _write_methodology,
    "summary": _write_summary,
    "ai_tools": _write_ai_tools,
    "ai_queries": _write_ai_queries,
    "file_lifecycle": _write_file_lifecycle,
    "exfil": _write_exfil,
    "attribution": _write_attribution,
    "timeline": _write_timeline,
    "exports": _write_exports,
    "coverage": _write_coverage,
    "limitations": _write_limitations,
    "custody": _write_custody,
    "signature": _write_signature,
}


def _start_section(doc: DocumentObject, *, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    wide = Emu(max(section.page_width or 0, section.page_height or 0))
    narrow = Emu(min(section.page_width or 0, section.page_height or 0))
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (wide, narrow) if landscape else (narrow, wide)


def render_docx(model: ReportModel, out_path: Path | str) -> None:
    lang = str(model.case.get("lang", DEFAULT_LANG))
    doc = Document()
    for name in _STYLES_WITH_FONT:
        _set_east_asian_font(doc.styles[name])
    first = doc.sections[0]
    first.left_margin = first.right_margin = Inches(PAGE_MARGIN_IN)
    doc.add_heading(t("report.title", lang), level=0)
    doc.add_paragraph(t("handling_notice", lang))
    for key in SECTION_KEYS:
        if key == "run":
            continue
        if key == "timeline":
            _start_section(doc, landscape=True)
        doc.add_heading(t(f"section.{key}", lang), level=1)
        _SECTION_WRITERS[key](doc, model, lang)
        if key == "timeline":
            _start_section(doc, landscape=False)
    doc.save(str(out_path))
