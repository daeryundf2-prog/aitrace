from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from aitrace.report.docx import EAST_ASIAN_FONT, MONO_FONT, render_docx
from aitrace.report.i18n import t
from aitrace.report.model import COLUMNS, ReportModel

KO_APPENDIX_E_HEADINGS = [
    "사건 정보",
    "증거물 및 무결성 검증",
    "분석 도구 및 환경",
    "분석 방법",
    "분석 결과 요약",
    "AI 도구 사용 현황",
    "AI 질의 및 응답",
    "파일 생성·수정·삭제 내역",
    "반출·원격·전송 경로",
    "사람/AI 작성 판단",
    "통합 타임라인",
    "추출 자료 목록",
    "분석 범위 및 미해석 항목",
    "한계 및 유의사항",
    "증거물 관리 이력",
    "서명",
]


def _populated(model: ReportModel) -> ReportModel:
    sha = "a1" * 32
    model.ai_tools = [
        {
            "kind_label": "CLI",
            "tool": "Claude Code",
            "users": ["kim", "lee"],
            "first_seen_local": "2026-01-01 09:00",
            "last_seen_local": "2026-01-02 09:00",
            "evidence_count": 3,
        }
    ]
    model.file_lifecycle = [
        {
            "severity_label": "높음",
            "display_title": "파일 a.docx",
            "display_steps": ["09:00 생성", "09:05 삭제"],
            "links_summary": "없음",
        }
    ]
    model.exfil = [
        {
            "severity_label": "높음",
            "display_channel": "Google Drive \u2192 drive.google.com",
            "ts_local": "09:10",
            "actor_label": "사람",
            "user": "kim",
            "source_name": "a.docx",
        }
    ]
    model.attribution = [
        {
            "subject": "a.docx",
            "verdict_label": "사람 작성",
            "confidence_display": "0.80 (높음)",
            "signals_summary": "신호 (지지)",
            "limitations_display": ["한계 A", "한계 B"],
        }
    ]
    model.timeline = [
        {
            "ts_local": "09:00",
            "category_label": "AI 활동",
            "actor_label": "AI",
            "user": "kim",
            "summary": "요약",
            "source_path": "/Users/kim/x",
        }
    ]
    model.exports = [{"export_path": "files/a.docx", "size": 5, "sha256": sha}]
    model.custody = [{"seq": 1, "actor": "QA", "action": "intake", "detail": {"hash": "ok"}}]
    return model


def _model(ai_queries: list[dict[str, object]] | None = None) -> ReportModel:
    return ReportModel(
        case={
            "number": "T-1",
            "name": "테스트 사건",
            "examiner": "QA",
            "organisation": "Lab",
            "timezone": "Asia/Seoul",
            "notes": "",
            "tool_version": "0.1.0",
            "lang": "ko",
            "handling_notice": "대외비 - 개인정보 포함 / Confidential - contains personal data",
        },
        evidence=[
            {
                "path": "/images/win.E01",
                "kind": "e01",
                "size": 10,
                "media_sha256": "bb",
                "verification_label": "일치",
            }
        ],
        tools={
            "aitrace_version": "0.1.0",
            "python": "3.12.12",
            "dissect": {"dissect.target": "3.25.1"},
            "os": "Darwin 24.6.0",
        },
        methodology={
            "paragraphs": ["증거물은 읽기 전용으로 열람하였다."],
            "standards_intro": "적용 기준",
            "standards": ["ISO/IEC 27037"],
            "tool_validation": "도구 검증: aitrace 0.1.0",
        },
        summary=[
            {"id": f"Q{i}", "question": f"Q{i}.", "bullets": ["해당 발견 사항 없음"]} for i in range(1, 6)
        ],
        ai_tools=[],
        ai_queries=ai_queries or [],
        file_lifecycle=[],
        exfil=[],
        attribution=[],
        timeline=[],
        exports=[],
        coverage={"plugin_runs": [], "parse_warnings": [], "not_decoded": [], "findings": []},
        limitations=["한계 항목"],
        custody=[{"seq": 1, "actor": "QA", "action": "intake", "detail": "hashed"}],
        signature={"examiner": "QA", "organisation": "Lab", "signature": "", "date": ""},
        run={},
    )


def test_docx_headings_match_appendix_e_for_ko(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    render_docx(_model(), out)

    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    assert headings == KO_APPENDIX_E_HEADINGS
    assert doc.styles["Normal"].element.rPr is not None
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == EAST_ASIAN_FONT
    header_rows = [table.rows[0]._tr.trPr for table in doc.tables if table.rows[0]._tr.trPr is not None]
    assert header_rows
    assert all(trpr.find(qn("w:tblHeader")) is not None for trpr in header_rows)


def test_docx_renders_very_long_cell_without_error(tmp_path: Path) -> None:
    long_text = "x" * 5000
    query: dict[str, object] = {
        "ts_local": None,
        "tool": "Claude Code",
        "role_label": "사용자",
        "text": long_text,
        "source_name": "s1.jsonl",
    }
    out = tmp_path / "long.docx"
    render_docx(_model(ai_queries=[query]), out)

    doc = Document(str(out))
    cell_lengths = [len(cell.text) for table in doc.tables for row in table.rows for cell in row.cells]
    assert max(cell_lengths) == len(long_text)


def _render_populated(tmp_path: Path) -> DocumentObject:
    out = tmp_path / "full.docx"
    render_docx(_populated(_model()), out)
    return Document(str(out))


def test_docx_spec_tables_use_columns_labels_in_ko(tmp_path: Path) -> None:
    doc = _render_populated(tmp_path)

    header_sets = [[c.text for c in table.rows[0].cells] for table in doc.tables]
    for section in ("ai_tools", "file_lifecycle", "exfil", "attribution", "timeline", "exports"):
        expected = [t(label_key, "ko") for _, label_key, _ in COLUMNS[section]]
        assert expected in header_sets, section


def test_docx_cells_never_contain_raw_json(tmp_path: Path) -> None:
    doc = _render_populated(tmp_path)

    texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert texts
    assert not [text for text in texts if text.startswith(("{", "["))]
    assert "hash: ok" in texts


def test_docx_column_widths_follow_weight_and_timeline_is_landscape(tmp_path: Path) -> None:
    doc = _render_populated(tmp_path)

    exports = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == t("col.export_path", "ko"))
    widths = [col.width for col in exports.columns]
    assert all(w is not None for w in widths)
    weights = [weight for _, _, weight in COLUMNS["exports"]]
    ratios = [
        (w or 0) * sum(weights) / (sum(w or 0 for w in widths) * weight)
        for w, weight in zip(widths, weights, strict=True)
    ]
    assert all(abs(r - 1) < 0.01 for r in ratios)

    landscape = [s for s in doc.sections if s.orientation == WD_ORIENT.LANDSCAPE]
    assert len(landscape) == 1
    assert (landscape[0].page_width or 0) > (landscape[0].page_height or 0)


def test_docx_hash_cells_use_small_monospace(tmp_path: Path) -> None:
    doc = _render_populated(tmp_path)

    exports = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == t("col.export_path", "ko"))
    run = exports.rows[1].cells[2].paragraphs[0].runs[0]
    assert run.font.name == MONO_FONT
    assert run.font.size is not None and run.font.size.pt < 9
    body_run = exports.rows[1].cells[0].paragraphs[0].runs[0]
    assert body_run.font.size is not None and body_run.font.size.pt == 9
